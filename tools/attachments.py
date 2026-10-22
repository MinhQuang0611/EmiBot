"""Save email attachments to disk and pull out their text, so the agent can read vouchers,
rooming lists and payment slips without opening Outlook.

PDF needs pypdf, Excel needs openpyxl (requirements.txt). Images are saved and listed; the
agent opens the file itself to read a payment slip screenshot. A scanned PDF has no text, so its
pages are rendered to PNG next to it (pypdfium2 + pillow) and the agent opens those instead.
"""
import csv
import hashlib
import io
import os
import re
import zipfile

MAX_TEXT = 20000           # characters of extracted text printed per file
SMALL_INLINE_IMAGE = 15000  # bytes; inline images under this are logos and signatures
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp", ".heic", ".tif", ".tiff"}
PAGE_DPI = 170              # scanned PDF pages: enough to read a voucher's small print
MAX_PAGE_IMAGES = 5         # scanned pages rendered per PDF


def message_dir(work_dir, message_id):
    """One folder per message, named by a short hash of the Graph id (ids are too long for paths)."""
    return os.path.join(work_dir, hashlib.sha1(message_id.encode()).hexdigest()[:12])


def safe_name(name):
    name = re.sub(r'[\\/:*?"<>|\x00-\x1f]', "_", name or "attachment").strip(" .")
    return name[:150] or "attachment"


def skip(att):
    """Inline logos and signature images are noise."""
    ext = os.path.splitext(att.get("name") or "")[1].lower()
    return att.get("isInline") and ext in IMAGE_EXT and len(att.get("data", b"")) < SMALL_INLINE_IMAGE


def save(folder, att):
    os.makedirs(folder, exist_ok=True)
    path = os.path.join(folder, safe_name(att.get("name")))
    with open(path, "wb") as f:
        f.write(att["data"])
    return path


def extract_text(path, data):
    """Best-effort text for the agent. Never raises: a broken file returns a note instead."""
    ext = os.path.splitext(path)[1].lower()
    try:
        if ext == ".pdf":
            return _pdf(path, data)
        if ext in (".xlsx", ".xlsm"):
            return _xlsx(data)
        if ext == ".docx":
            return _docx(data)
        if ext in (".csv", ".txt"):
            return _decode(data)
        if ext in IMAGE_EXT:
            return "(image — open the file to read it)"
        if ext == ".xls":
            return "(old .xls format, not extracted — ask for .xlsx or read it another way)"
        return f"({ext or 'no extension'} not extracted)"
    except ImportError as e:
        return f"(not extracted: {e.name} is not installed — pip install -r requirements.txt)"
    except Exception as e:
        return f"(could not read file: {type(e).__name__}: {e})"


def _decode(data):
    for enc in ("utf-8-sig", "cp1258", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", "replace")


def _pdf(path, data):
    from pypdf import PdfReader
    pages = [p.extract_text() or "" for p in PdfReader(io.BytesIO(data)).pages]
    scanned = [i for i, t in enumerate(pages, 1) if not t.strip()]
    if pages and len(scanned) == len(pages):
        # 2026-10-22: an MG voucher scan printed only empty page markers and the row went in blank
        return _page_images(path, data, scanned, "PDF has no text layer (scan).")
    text = "\n".join(f"--- page {i} ---\n{t}" for i, t in enumerate(pages, 1))
    if scanned:
        text += "\n\n" + _page_images(path, data, scanned,
                                      f"PDF page(s) {', '.join(map(str, scanned))} have no text layer (scan).")
    return text


def _page_images(path, data, numbers, head):
    """Render scanned pages to <name>-pageN.png next to the PDF, so the agent can open them like any image."""
    try:
        import pypdfium2 as pdfium
        doc = pdfium.PdfDocument(data)
        try:
            saved = []
            for n in numbers[:MAX_PAGE_IMAGES]:
                png = f"{os.path.splitext(path)[0]}-page{n}.png"
                doc[n - 1].render(scale=PAGE_DPI / 72).to_pil().save(png)
                saved.append(os.path.abspath(png))
        finally:
            doc.close()
    except ImportError as e:
        pkg = {"PIL": "pillow"}.get(e.name, e.name)
        return f"{head} Page images not made: {pkg} is not installed — pip install {pkg}, then run show again."
    except Exception as e:
        return f"{head} Page images not made ({type(e).__name__}: {e}) — open the PDF another way."
    more = (f"\nOnly the first {MAX_PAGE_IMAGES} of {len(numbers)} scanned pages were saved."
            if len(numbers) > MAX_PAGE_IMAGES else "")
    return f"{head} Page images saved — open each one to read it:\n" + "\n".join(saved) + more


def _xlsx(data):
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"--- sheet {ws.title} ---")
        buf = io.StringIO()
        w = csv.writer(buf, delimiter="\t", lineterminator="\n")
        for row in ws.iter_rows(values_only=True):
            if any(c not in (None, "") for c in row):
                w.writerow(["" if c is None else c for c in row])
        out.append(buf.getvalue())
    return "\n".join(out)


def _docx(data):
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        xml = z.read("word/document.xml").decode("utf-8")
    xml = re.sub(r"</w:p>", "\n", xml)
    return re.sub(r"<[^>]+>", "", xml)


def clip(text):
    return text if len(text) <= MAX_TEXT else text[:MAX_TEXT] + f"\n... ({len(text) - MAX_TEXT} more characters in the file)"
